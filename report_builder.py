"""Builds the final Excel report: a Players sheet with color-coded
upcoming fixtures (matching the official FPL app's difficulty colors),
and a Price Changes sheet."""

import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Approximate the official FPL app's Fixture Difficulty Rating colors.
FDR_FILL = {
    1: "375F2D",  # very easy - dark green
    2: "01FC7A",  # easy - green
    3: "E7E7E7",  # medium - grey
    4: "FF1751",  # hard - red
    5: "80072D",  # very hard - dark red
}
FDR_FONT = {
    1: "FFFFFF",
    2: "000000",
    3: "000000",
    4: "FFFFFF",
    5: "FFFFFF",
}

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _style_header(ws, headers):
    for col_idx, name in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _autofit(ws, headers):
    for col_idx, name in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(12, len(name) + 4)


def write_players_sheet(ws, df: pd.DataFrame, num_fixtures: int):
    display_cols = ["Player", "Team", "Position", "Price", "Form", "Selected By %", "Total Points"]
    display_cols += [f"Fixture {i + 1}" for i in range(num_fixtures)]

    _style_header(ws, display_cols)

    for row_idx, (_, row) in enumerate(df.iterrows(), start=2):
        for col_idx, col_name in enumerate(display_cols, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row[col_name])

            if col_name.startswith("Fixture"):
                difficulty = row.get(f"{col_name} FDR")
                if pd.notna(difficulty):
                    difficulty = int(difficulty)
                    cell.fill = PatternFill(
                        start_color=FDR_FILL.get(difficulty, "FFFFFF"),
                        end_color=FDR_FILL.get(difficulty, "FFFFFF"),
                        fill_type="solid",
                    )
                    cell.font = Font(color=FDR_FONT.get(difficulty, "000000"), bold=True)
                    cell.alignment = Alignment(horizontal="center")

    _autofit(ws, display_cols)
    ws.freeze_panes = "A2"


def write_price_changes_sheet(ws, changes_df: pd.DataFrame):
    if changes_df.empty:
        ws.cell(
            row=1, column=1,
            value="No previous snapshot to compare yet — run this again tomorrow to see price changes.",
        )
        return

    headers = list(changes_df.columns)
    _style_header(ws, headers)

    for row_idx, (_, row) in enumerate(changes_df.iterrows(), start=2):
        for col_idx, col_name in enumerate(headers, start=1):
            value = row[col_name]
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if col_name == "Change":
                cell.font = Font(color="1B5E20" if value > 0 else "B71C1C", bold=True)

    _autofit(ws, headers)


def build_report(df: pd.DataFrame, changes_df: pd.DataFrame, output_file: str, num_fixtures: int):
    os.makedirs(os.path.dirname(output_file), exist_ok=True)

    wb = Workbook()
    ws_players = wb.active
    ws_players.title = "Players"
    write_players_sheet(ws_players, df, num_fixtures)

    ws_changes = wb.create_sheet("Price Changes")
    write_price_changes_sheet(ws_changes, changes_df)

    wb.save(output_file)
